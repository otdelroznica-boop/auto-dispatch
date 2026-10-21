import io
import json
import secrets
from datetime import date as date_cls, timedelta
from pathlib import Path
from typing import Optional

from fastapi import Depends, FastAPI, HTTPException, Request, Response
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles

from . import telegram
from .database import check_pin, get_connection, hash_pin, init_db, normalize_phone
from .models import (ClientIn, DriverIn, LoginIn, LostIn, OrderIn, PayIn, ReasonIn, ServiceIn, SettingsIn, StatusIn,
                     UserIn, VehicleIn)

STATIC_DIR = Path(__file__).resolve().parent / "static"
COOKIE = "dispatch_session"

app = FastAPI(title="Диспетчерская")


@app.on_event("startup")
def startup_event():
    init_db()


def db():
    conn = get_connection()
    try:
        yield conn
    finally:
        conn.close()


def rows(cursor) -> list[dict]:
    return [dict(r) for r in cursor.fetchall()]


def settings_dict(conn) -> dict:
    return {r["key"]: r["value"] for r in conn.execute("SELECT key, value FROM settings")}


# ---------- Авторизация ----------

def current_user(request: Request, conn=Depends(db)) -> dict:
    token = request.cookies.get(COOKIE)
    if token:
        row = conn.execute(
            "SELECT u.id, u.name, u.is_admin FROM sessions s JOIN users u ON u.id = s.user_id "
            "WHERE s.token = ? AND u.active = 1", (token,)).fetchone()
        if row:
            return dict(row)
    raise HTTPException(401, "Нужно войти")


def admin_user(user=Depends(current_user)) -> dict:
    if not user["is_admin"]:
        raise HTTPException(403, "Доступно только администратору")
    return user


@app.get("/api/login-users")
def login_users(conn=Depends(db)):
    return rows(conn.execute("SELECT id, name FROM users WHERE active = 1 ORDER BY name"))


@app.post("/api/login")
def login(data: LoginIn, response: Response, conn=Depends(db)):
    row = conn.execute("SELECT id, pin_hash FROM users WHERE id = ? AND active = 1", (data.user_id,)).fetchone()
    if not row or not check_pin(data.pin, row["pin_hash"]):
        raise HTTPException(401, "Неверный пин-код")
    token = secrets.token_urlsafe(32)
    conn.execute("INSERT INTO sessions(token, user_id) VALUES (?, ?)", (token, row["id"]))
    conn.commit()
    response.set_cookie(COOKIE, token, max_age=60 * 60 * 24 * 180, httponly=True, samesite="lax")
    return {"ok": True}


@app.post("/api/logout")
def logout(request: Request, response: Response, conn=Depends(db)):
    conn.execute("DELETE FROM sessions WHERE token = ?", (request.cookies.get(COOKIE, ""),))
    conn.commit()
    response.delete_cookie(COOKIE)
    return {"ok": True}


# ---------- Справочники ----------

@app.get("/api/bootstrap")
def bootstrap(user=Depends(current_user), conn=Depends(db)):
    s = settings_dict(conn)
    if not user["is_admin"]:
        s.pop("tg_token", None)
    else:
        s["tg_token_set"] = bool(s.get("tg_token"))
        s.pop("tg_token", None)
    return {
        "me": user,
        "vehicles": rows(conn.execute("SELECT * FROM vehicles ORDER BY sort, id")),
        "drivers": rows(conn.execute("SELECT * FROM drivers ORDER BY name")),
        "services": rows(conn.execute("SELECT * FROM services ORDER BY vehicle_id IS NULL, vehicle_id, id")),
        "reasons": rows(conn.execute("SELECT * FROM lost_reasons ORDER BY sort, id")),
        "settings": s,
    }


def upsert(conn, table: str, item_id: Optional[int], data: dict) -> int:
    if item_id:
        sets = ", ".join(f"{k} = ?" for k in data)
        conn.execute(f"UPDATE {table} SET {sets} WHERE id = ?", (*data.values(), item_id))
    else:
        cols = ", ".join(data)
        marks = ", ".join("?" for _ in data)
        item_id = conn.execute(f"INSERT INTO {table}({cols}) VALUES ({marks})", tuple(data.values())).lastrowid
    conn.commit()
    return item_id


@app.post("/api/vehicles")
@app.put("/api/vehicles/{item_id}")
def save_vehicle(data: VehicleIn, item_id: Optional[int] = None, user=Depends(current_user), conn=Depends(db)):
    return {"id": upsert(conn, "vehicles", item_id, data.model_dump())}


@app.post("/api/drivers")
@app.put("/api/drivers/{item_id}")
def save_driver(data: DriverIn, item_id: Optional[int] = None, user=Depends(current_user), conn=Depends(db)):
    return {"id": upsert(conn, "drivers", item_id, data.model_dump())}


@app.post("/api/services")
@app.put("/api/services/{item_id}")
def save_service(data: ServiceIn, item_id: Optional[int] = None, user=Depends(current_user), conn=Depends(db)):
    return {"id": upsert(conn, "services", item_id, data.model_dump())}


@app.post("/api/reasons")
@app.put("/api/reasons/{item_id}")
def save_reason(data: ReasonIn, item_id: Optional[int] = None, user=Depends(current_user), conn=Depends(db)):
    return {"id": upsert(conn, "lost_reasons", item_id, data.model_dump())}


# ---------- Отказы (упущенная прибыль) ----------

LOST_SQL = """
SELECT l.*, v.name AS vehicle_name, v.color AS vehicle_color, r.title AS reason_title, u.name AS created_by_name
FROM lost_orders l
LEFT JOIN vehicles v ON v.id = l.vehicle_id
LEFT JOIN lost_reasons r ON r.id = l.reason_id
LEFT JOIN users u ON u.id = l.created_by
"""


def lost_row(row) -> dict:
    item = dict(row)
    item["lines"] = json.loads(item.pop("lines_json") or "[]")
    return item


def load_lost(conn, date_from: str, date_to: str) -> list[dict]:
    return [lost_row(r) for r in conn.execute(LOST_SQL + " WHERE l.date BETWEEN ? AND ? ORDER BY l.date, l.time, l.id",
                                              (date_from, date_to))]


@app.get("/api/lost")
def list_lost(date: Optional[str] = None, date_from: Optional[str] = None, date_to: Optional[str] = None,
              user=Depends(current_user), conn=Depends(db)):
    date_from = date_from or date or date_cls.today().isoformat()
    return load_lost(conn, date_from, date_to or date or date_from)


@app.post("/api/lost")
@app.put("/api/lost/{item_id}")
def save_lost(data: LostIn, item_id: Optional[int] = None, user=Depends(current_user), conn=Depends(db)):
    lines = [l.model_dump() for l in data.lines if l.title.strip()]
    payload = {
        "date": data.date, "time": data.time, "vehicle_id": data.vehicle_id, "reason_id": data.reason_id,
        "lines_json": json.dumps(lines, ensure_ascii=False),
        "amount": round(sum(l["qty"] * l["price"] for l in lines), 2) if lines else data.amount,
        "client_phone": data.client_phone, "note": data.note, "partner": int(data.partner),
        "partner_name": data.partner_name if data.partner else "",
        "partner_income": data.partner_income if data.partner else 0,
    }
    if not item_id:
        payload["created_by"] = user["id"]
    item_id = upsert(conn, "lost_orders", item_id, payload)
    return lost_row(conn.execute(LOST_SQL + " WHERE l.id = ?", (item_id,)).fetchone())


@app.delete("/api/lost/{item_id}")
def delete_lost(item_id: int, user=Depends(current_user), conn=Depends(db)):
    conn.execute("DELETE FROM lost_orders WHERE id = ?", (item_id,))
    conn.commit()
    return {"ok": True}


# ---------- Клиенты ----------

@app.get("/api/clients")
def search_clients(q: str = "", limit: int = 20, user=Depends(current_user), conn=Depends(db)):
    digits = "".join(ch for ch in q if ch.isdigit())
    base = ("SELECT c.*, COUNT(o.id) AS orders_count, COALESCE(SUM(o.paid_amount), 0) AS paid_total, "
            "MAX(o.date) AS last_date FROM clients c "
            "LEFT JOIN orders o ON o.client_id = c.id AND o.status != 'cancel' ")
    if len(digits) >= 3:
        # 8 и +7 в начале равнозначны: ищем по хвосту номера
        tail = digits[1:] if len(digits) >= 4 and digits[0] in "78" else digits
        cur = conn.execute(base + "WHERE c.phone_norm LIKE ? GROUP BY c.id ORDER BY last_date DESC LIMIT ?",
                           (f"%{tail}%", limit))
    elif q.strip():
        cur = conn.execute(base + "WHERE c.name LIKE ? OR c.name LIKE ? OR c.inn LIKE ? "
                           "GROUP BY c.id ORDER BY last_date DESC LIMIT ?",
                           (f"%{q.strip()}%", f"%{q.strip().capitalize()}%", f"%{q.strip()}%", limit))
    else:
        cur = conn.execute(base + "GROUP BY c.id ORDER BY last_date DESC, c.id DESC LIMIT ?", (limit,))
    return rows(cur)


@app.get("/api/clients/{client_id}")
def get_client(client_id: int, user=Depends(current_user), conn=Depends(db)):
    row = conn.execute("SELECT * FROM clients WHERE id = ?", (client_id,)).fetchone()
    if not row:
        raise HTTPException(404, "Клиент не найден")
    orders = [order_summary(r) for r in conn.execute(ORDER_SQL + " WHERE o.client_id = ? ORDER BY o.date DESC, "
                                                                  "o.start_time DESC LIMIT 50", (client_id,))]
    return {**dict(row), "orders": orders}


@app.post("/api/clients")
@app.put("/api/clients/{client_id}")
def save_client(data: ClientIn, client_id: Optional[int] = None, user=Depends(current_user), conn=Depends(db)):
    payload = data.model_dump()
    payload["phone_norm"] = normalize_phone(data.phone)
    if payload["phone_norm"]:
        dup = conn.execute("SELECT id FROM clients WHERE phone_norm = ? AND id != ?",
                           (payload["phone_norm"], client_id or 0)).fetchone()
        if dup:
            raise HTTPException(409, "Клиент с таким телефоном уже есть")
    return {"id": upsert(conn, "clients", client_id, payload)}


def resolve_client(conn, client: Optional[ClientIn]) -> Optional[int]:
    """Найти клиента по id или телефону; если нет — создать. Пустой клиент допустим."""
    if client is None:
        return None
    phone_norm = normalize_phone(client.phone)
    if client.id:
        conn.execute("UPDATE clients SET kind = ?, name = ?, phone = ?, phone_norm = ?, inn = ? WHERE id = ?",
                     (client.kind, client.name.strip(), client.phone, phone_norm, client.inn, client.id))
        return client.id
    if phone_norm:
        row = conn.execute("SELECT id, name FROM clients WHERE phone_norm = ?", (phone_norm,)).fetchone()
        if row:
            if client.name.strip() and not row["name"]:
                conn.execute("UPDATE clients SET name = ? WHERE id = ?", (client.name.strip(), row["id"]))
            return row["id"]
    if not phone_norm and not client.name.strip():
        return None
    return conn.execute(
        "INSERT INTO clients(kind, name, phone, phone_norm, inn) VALUES (?, ?, ?, ?, ?)",
        (client.kind, client.name.strip(), client.phone, phone_norm, client.inn)).lastrowid


# ---------- Заказы ----------

ORDER_SQL = """
SELECT o.*, v.name AS vehicle_name, v.color AS vehicle_color,
       d.name AS driver_name, d.phone AS driver_phone, d.tg_chat_id AS driver_chat,
       c.name AS client_name, c.phone AS client_phone, c.kind AS client_kind, c.inn AS client_inn,
       u.name AS created_by_name
FROM orders o
JOIN vehicles v ON v.id = o.vehicle_id
LEFT JOIN drivers d ON d.id = o.driver_id
LEFT JOIN clients c ON c.id = o.client_id
LEFT JOIN users u ON u.id = o.created_by
"""


def order_summary(row) -> dict:
    o = dict(row)
    o.pop("driver_chat", None)
    return o


def load_order(conn, order_id: int) -> dict:
    row = conn.execute(ORDER_SQL + " WHERE o.id = ?", (order_id,)).fetchone()
    if not row:
        raise HTTPException(404, "Заказ не найден")
    order = dict(row)
    order["lines"] = rows(conn.execute("SELECT id, service_id, title, qty, price FROM order_lines "
                                       "WHERE order_id = ? ORDER BY id", (order_id,)))
    return order


def attach_lines(conn, orders: list[dict]) -> list[dict]:
    if not orders:
        return orders
    ids = [o["id"] for o in orders]
    marks = ",".join("?" for _ in ids)
    by_order: dict[int, list] = {i: [] for i in ids}
    for line in conn.execute(f"SELECT order_id, service_id, title, qty, price FROM order_lines "
                             f"WHERE order_id IN ({marks}) ORDER BY id", ids):
        by_order[line["order_id"]].append(dict(line))
    for o in orders:
        o["lines"] = by_order[o["id"]]
    return orders


@app.get("/api/orders")
def list_orders(date: Optional[str] = None, date_from: Optional[str] = None, date_to: Optional[str] = None,
                user=Depends(current_user), conn=Depends(db)):
    date_from = date_from or date or date_cls.today().isoformat()
    date_to = date_to or date or date_from
    cur = conn.execute(ORDER_SQL + " WHERE o.date BETWEEN ? AND ? ORDER BY o.date, o.start_time",
                       (date_from, date_to))
    return attach_lines(conn, [order_summary(r) for r in cur])


@app.get("/api/orders/{order_id}")
def get_order(order_id: int, user=Depends(current_user), conn=Depends(db)):
    return order_summary(load_order(conn, order_id))


def find_conflict(conn, data: OrderIn, order_id: Optional[int]) -> Optional[dict]:
    row = conn.execute(
        ORDER_SQL + " WHERE o.vehicle_id = ? AND o.date = ? AND o.status != 'cancel' AND o.id != ? "
                    "AND o.start_time < ? AND o.end_time > ? LIMIT 1",
        (data.vehicle_id, data.date, order_id or 0, data.end_time, data.start_time)).fetchone()
    return dict(row) if row else None


def notify(conn, order: dict, header: str):
    s = settings_dict(conn)
    telegram.send(s.get("tg_token", ""), [order.get("driver_chat") or "", s.get("tg_dispatch_chat", "")],
                  telegram.order_text(order, header))


@app.post("/api/orders")
@app.put("/api/orders/{order_id}")
def save_order(data: OrderIn, order_id: Optional[int] = None, user=Depends(current_user), conn=Depends(db)):
    if data.end_time <= data.start_time:
        raise HTTPException(400, "Время окончания должно быть позже начала")
    if not data.force:
        conflict = find_conflict(conn, data, order_id)
        if conflict:
            raise HTTPException(409, {
                "message": f"Машина «{conflict['vehicle_name']}» уже занята с {conflict['start_time']} до "
                           f"{conflict['end_time']} (заказ #{conflict['id']}, {conflict['client_name'] or 'без имени'})",
                "conflict_id": conflict["id"]})

    before = load_order(conn, order_id) if order_id else None
    client_id = resolve_client(conn, data.client)
    lines = [l for l in data.lines if l.title.strip()]
    plan = round(sum(l.qty * l.price for l in lines), 2) if lines else data.plan_amount
    fields = {
        "vehicle_id": data.vehicle_id, "driver_id": data.driver_id, "client_id": client_id,
        "date": data.date, "start_time": data.start_time, "end_time": data.end_time,
        "address_from": data.address_from, "address_to": data.address_to, "note": data.note,
        "status": data.status, "plan_amount": plan, "paid_amount": data.paid_amount,
        "payment_method": data.payment_method,
    }
    if order_id:
        sets = ", ".join(f"{k} = ?" for k in fields)
        conn.execute(f"UPDATE orders SET {sets}, updated_at = CURRENT_TIMESTAMP WHERE id = ?",
                     (*fields.values(), order_id))
        conn.execute("DELETE FROM order_lines WHERE order_id = ?", (order_id,))
    else:
        fields["created_by"] = user["id"]
        cols = ", ".join(fields)
        order_id = conn.execute(f"INSERT INTO orders({cols}) VALUES ({', '.join('?' for _ in fields)})",
                                tuple(fields.values())).lastrowid
    conn.executemany("INSERT INTO order_lines(order_id, service_id, title, qty, price) VALUES (?, ?, ?, ?, ?)",
                     [(order_id, l.service_id, l.title.strip(), l.qty, l.price) for l in lines])
    conn.commit()

    order = load_order(conn, order_id)
    if before is None:
        notify(conn, order, "Новый заказ")
    elif data.status == "cancel" and before["status"] != "cancel":
        notify(conn, order, "Заказ отменён")
    elif any(before[k] != order[k] for k in ("date", "start_time", "end_time", "vehicle_id", "driver_id",
                                              "address_from", "address_to")):
        notify(conn, order, "Заказ изменён")
    return order_summary(order)


@app.patch("/api/orders/{order_id}/status")
def set_status(order_id: int, data: StatusIn, user=Depends(current_user), conn=Depends(db)):
    before = load_order(conn, order_id)
    conn.execute("UPDATE orders SET status = ?, updated_at = CURRENT_TIMESTAMP WHERE id = ?", (data.status, order_id))
    conn.commit()
    order = load_order(conn, order_id)
    if data.status == "cancel" and before["status"] != "cancel":
        notify(conn, order, "Заказ отменён")
    return order_summary(order)


@app.patch("/api/orders/{order_id}/pay")
def set_paid(order_id: int, data: PayIn, user=Depends(current_user), conn=Depends(db)):
    load_order(conn, order_id)
    conn.execute("UPDATE orders SET paid_amount = ?, payment_method = COALESCE(?, payment_method), "
                 "status = CASE WHEN ? THEN 'done' ELSE status END, updated_at = CURRENT_TIMESTAMP WHERE id = ?",
                 (data.paid_amount, data.payment_method, 1 if data.mark_done else 0, order_id))
    conn.commit()
    return order_summary(load_order(conn, order_id))


@app.delete("/api/orders/{order_id}")
def delete_order(order_id: int, user=Depends(current_user), conn=Depends(db)):
    conn.execute("DELETE FROM orders WHERE id = ?", (order_id,))
    conn.commit()
    return {"ok": True}


# ---------- Отчёты ----------

def days_between(date_from: str, date_to: str) -> list[str]:
    d, end = date_cls.fromisoformat(date_from), date_cls.fromisoformat(date_to)
    out = []
    while d <= end and len(out) < 400:
        out.append(d.isoformat())
        d += timedelta(days=1)
    return out


def debt_of(o: dict) -> float:
    """Долг — только по выполненным заказам: будущие ещё не оплачены, но и не должны."""
    return max(o["plan_amount"] - o["paid_amount"], 0) if o["status"] == "done" else 0


@app.get("/api/report")
def report(date_from: str, date_to: str, user=Depends(current_user), conn=Depends(db)):
    orders = attach_lines(conn, [order_summary(r) for r in conn.execute(
        ORDER_SQL + " WHERE o.date BETWEEN ? AND ? ORDER BY o.date, o.start_time", (date_from, date_to))])
    active = [o for o in orders if o["status"] != "cancel"]

    def group(key_fn, name_fn):
        out: dict = {}
        for o in active:
            key = key_fn(o)
            g = out.setdefault(key, {"id": key, "name": name_fn(o), "orders": 0, "plan": 0.0, "paid": 0.0, "debt": 0.0})
            g["orders"] += 1
            g["plan"] += o["plan_amount"]
            g["paid"] += o["paid_amount"]
            g["debt"] += debt_of(o)
        return sorted(out.values(), key=lambda g: -g["paid"])

    services: dict = {}
    for o in active:
        for l in o["lines"]:
            g = services.setdefault(l["title"], {"name": l["title"], "qty": 0.0, "amount": 0.0, "orders": 0})
            g["qty"] += l["qty"]
            g["amount"] += l["qty"] * l["price"]
            g["orders"] += 1

    by_day: dict = {}
    for o in active:
        g = by_day.setdefault(o["date"], {"date": o["date"], "orders": 0, "plan": 0.0, "paid": 0.0})
        g["orders"] += 1
        g["plan"] += o["plan_amount"]
        g["paid"] += o["paid_amount"]

    debts = [o for o in active if debt_of(o) > 0.009]
    return {
        "date_from": date_from, "date_to": date_to,
        "totals": {
            "orders": len(active),
            "cancelled": len(orders) - len(active),
            "plan": sum(o["plan_amount"] for o in active),
            "paid": sum(o["paid_amount"] for o in active),
            "debt": sum(debt_of(o) for o in active),
        },
        "by_vehicle": group(lambda o: o["vehicle_id"], lambda o: o["vehicle_name"]),
        "by_driver": group(lambda o: o["driver_id"] or 0, lambda o: o["driver_name"] or "Не указан"),
        "by_payment": group(lambda o: o["payment_method"],
                            lambda o: telegram.PAY_TEXT.get(o["payment_method"], o["payment_method"]).capitalize()),
        "by_service": sorted(services.values(), key=lambda g: -g["amount"]),
        "by_day": [by_day.get(d) or {"date": d, "orders": 0, "plan": 0.0, "paid": 0.0} for d in days_between(date_from, date_to)],
        "debts": debts,
        "lost": load_lost(conn, date_from, date_to),
    }


@app.get("/api/export")
def export(date_from: str, date_to: str, user=Depends(current_user), conn=Depends(db)):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    orders = attach_lines(conn, [order_summary(r) for r in conn.execute(
        ORDER_SQL + " WHERE o.date BETWEEN ? AND ? ORDER BY o.date, o.start_time", (date_from, date_to))])
    wb = Workbook()
    ws = wb.active
    ws.title = "Заказы"
    head = ["№", "Дата", "С", "По", "Машина", "Водитель", "Клиент", "Телефон", "Откуда", "Куда", "Услуги",
            "Расчёт, ₽", "Взяли, ₽", "Долг, ₽", "Оплата", "Статус", "Комментарий"]
    ws.append(head)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for o in orders:
        services = "; ".join(l["title"] + (f" × {l['qty']:g}" if l["qty"] != 1 else "") for l in o["lines"])
        debt = debt_of(o)
        ws.append([o["id"], o["date"], o["start_time"], o["end_time"], o["vehicle_name"], o["driver_name"] or "",
                   o["client_name"] or "", o["client_phone"] or "", o["address_from"], o["address_to"], services,
                   o["plan_amount"], o["paid_amount"], debt, telegram.PAY_TEXT.get(o["payment_method"], ""),
                   telegram.STATUS_TEXT.get(o["status"], o["status"]), o["note"]])
    for col, width in zip("ABCDEFGHIJKLMNOPQ", [6, 11, 7, 7, 13, 14, 22, 17, 28, 28, 40, 11, 11, 10, 12, 11, 30]):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"

    # Отказы
    ws2 = wb.create_sheet("Отказы")
    ws2.append(["Дата", "Время", "Машина", "Причина", "Услуги", "Упущено, ₽", "Передали партнёру", "Получили с передачи, ₽",
                "Телефон", "Комментарий"])
    for cell in ws2[1]:
        cell.font = Font(bold=True)
    for l in load_lost(conn, date_from, date_to):
        ws2.append([l["date"], l["time"], l["vehicle_name"] or "", l["reason_title"] or "",
                    "; ".join(x["title"] + (f" × {x['qty']:g}" if x["qty"] != 1 else "") for x in l["lines"]),
                    l["amount"], l["partner_name"] or "да" if l["partner"] else "", l["partner_income"] if l["partner"] else 0,
                    l["client_phone"], l["note"]])
    for col, width in zip("ABCDEFGHIJ", [11, 7, 14, 24, 40, 12, 18, 14, 17, 30]):
        ws2.column_dimensions[col].width = width
    ws2.freeze_panes = "A2"

    # Наёмные машины: сколько отдать владельцу
    hired = {r["id"]: dict(r) for r in conn.execute("SELECT * FROM vehicles WHERE ownership = 'hired'")}
    if hired:
        ws3 = wb.create_sheet("Наёмные машины")
        ws3.append(["Машина", "Владелец", "Заказов", "Взяли, ₽", "Отдаём владельцу, %", "Владельцу, ₽", "Остаётся нам, ₽"])
        for cell in ws3[1]:
            cell.font = Font(bold=True)
        for vid, v in hired.items():
            mine = [o for o in orders if o["vehicle_id"] == vid and o["status"] != "cancel"]
            paid = sum(o["paid_amount"] for o in mine)
            owner = round(paid * (v["owner_share"] or 0) / 100, 2)
            ws3.append([v["name"], v["owner_name"], len(mine), paid, v["owner_share"], owner, paid - owner])
        for col, width in zip("ABCDEFG", [16, 18, 9, 12, 18, 14, 16]):
            ws3.column_dimensions[col].width = width

    wb.properties.creator = ""
    wb.properties.lastModifiedBy = ""

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    name = f"orders_{date_from}_{date_to}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{name}"'})


# ---------- Настройки (админ) ----------

@app.put("/api/settings")
def save_settings(data: SettingsIn, user=Depends(admin_user), conn=Depends(db)):
    for key, value in data.model_dump(exclude_none=True).items():
        conn.execute("INSERT INTO settings(key, value) VALUES (?, ?) ON CONFLICT(key) DO UPDATE SET value = ?",
                     (key, str(value), str(value)))
    conn.commit()
    return {"ok": True}


@app.get("/api/telegram/chats")
def telegram_chats(user=Depends(admin_user), conn=Depends(db)):
    token = settings_dict(conn).get("tg_token")
    if not token:
        raise HTTPException(400, "Сначала сохраните токен бота")
    try:
        return telegram.discover_chats(token)
    except Exception as exc:
        raise HTTPException(502, f"Telegram не ответил: {exc}")


@app.post("/api/telegram/test")
def telegram_test(chat_id: str, user=Depends(admin_user), conn=Depends(db)):
    token = settings_dict(conn).get("tg_token")
    if not token:
        raise HTTPException(400, "Сначала сохраните токен бота")
    try:
        telegram._call(token, "sendMessage", {"chat_id": chat_id, "text": "Проверка связи: диспетчерская на связи."})
    except Exception as exc:
        raise HTTPException(502, f"Не отправилось: {exc}")
    return {"ok": True}


@app.get("/api/users")
def list_users(user=Depends(admin_user), conn=Depends(db)):
    return rows(conn.execute("SELECT id, name, is_admin, active FROM users ORDER BY id"))


@app.post("/api/users")
@app.put("/api/users/{item_id}")
def save_user(data: UserIn, item_id: Optional[int] = None, user=Depends(admin_user), conn=Depends(db)):
    payload = {"name": data.name, "is_admin": int(data.is_admin), "active": int(data.active)}
    if data.pin:
        if not data.pin.isdigit() or len(data.pin) != 4:
            raise HTTPException(400, "Пин-код — 4 цифры")
        payload["pin_hash"] = hash_pin(data.pin)
    elif not item_id:
        raise HTTPException(400, "Задайте пин-код")
    if item_id == user["id"] and (not data.is_admin or not data.active):
        raise HTTPException(400, "Нельзя снять права администратора с себя")
    return {"id": upsert(conn, "users", item_id, payload)}


# ---------- Статика ----------

app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")


@app.get("/sw.js")
def service_worker():
    return FileResponse(STATIC_DIR / "sw.js", media_type="application/javascript",
                        headers={"Cache-Control": "no-cache"})


@app.get("/manifest.webmanifest")
def manifest():
    return FileResponse(STATIC_DIR / "manifest.webmanifest", media_type="application/manifest+json")


@app.get("/")
def index():
    return FileResponse(STATIC_DIR / "index.html", headers={"Cache-Control": "no-cache"})
